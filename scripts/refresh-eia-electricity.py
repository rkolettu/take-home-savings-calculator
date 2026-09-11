from __future__ import annotations

import io
import json
import math
import re
import statistics
import sys
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

COST_DATA_PATH = Path(__file__).resolve().parents[1] / "src" / "data" / "sourcedCosts.json"
EIA_BILL_URL = "https://www.eia.gov/electricity/sales_revenue_price/xls/table_5A.xlsx"

STATE_CODES = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR",
    "California": "CA", "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE",
    "District of Columbia": "DC", "Florida": "FL", "Georgia": "GA", "Hawaii": "HI",
    "Idaho": "ID", "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS",
    "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN", "Mississippi": "MS",
    "Missouri": "MO", "Montana": "MT", "Nebraska": "NE", "Nevada": "NV",
    "New Hampshire": "NH", "New Jersey": "NJ", "New Mexico": "NM", "New York": "NY",
    "North Carolina": "NC", "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK",
    "Oregon": "OR", "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
    "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT",
    "Vermont": "VT", "Virginia": "VA", "Washington": "WA", "West Virginia": "WV",
    "Wisconsin": "WI", "Wyoming": "WY",
}


def normalize(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def state_code(value: object) -> str | None:
    text = normalize(value)
    if not text:
        return None
    # EIA occasionally appends footnote markers or superscript-style numbers.
    text = re.sub(r"^[^A-Za-z]+", "", text)
    text = re.sub(r"[^A-Za-z]+$", "", text)
    text = re.sub(r"\s+\d+$", "", text).strip()
    if text in STATE_CODES:
        return STATE_CODES[text]
    for name, code in STATE_CODES.items():
        if text.startswith(f"{name} "):
            return code
    return None


def number(value: object) -> float | None:
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = normalize(value).replace(",", "").replace("$", "")
        text = re.sub(r"[^0-9.\-]", "", text)
        if not text or text in {"-", ".", "-."}:
            return None
        try:
            result = float(text)
        except ValueError:
            return None
    return result if math.isfinite(result) else None


def download(url: str) -> bytes:
    request = urllib.request.Request(
        url,
        headers={"User-Agent": "take-home-savings-calculator data refresh"},
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        if response.status != 200:
            raise RuntimeError(f"EIA download returned HTTP {response.status}")
        return response.read()


def find_state_column(rows: list[list[object]]) -> int:
    width = max((len(row) for row in rows), default=0)
    counts: list[tuple[int, int]] = []
    for index in range(width):
        count = sum(
            1 for row in rows if index < len(row) and state_code(row[index]) is not None
        )
        counts.append((count, index))
    count, index = max(counts, default=(0, -1))
    if count < 45:
        raise RuntimeError(f"Could not reliably identify EIA state column ({count} matches)")
    return index


def find_bill_column(rows: list[list[object]], state_col: int) -> int:
    width = max((len(row) for row in rows), default=0)

    # First prefer a column whose visible header says monthly bill.
    header_candidates: set[int] = set()
    for row in rows[:25]:
        for index, value in enumerate(row):
            text = normalize(value).lower()
            if "monthly bill" in text or ("bill" in text and "dollar" in text):
                header_candidates.add(index)

    def candidate_values(index: int) -> list[float]:
        values: list[float] = []
        for row in rows:
            if state_col >= len(row) or index >= len(row):
                continue
            if state_code(row[state_col]) is None:
                continue
            value = number(row[index])
            if value is not None and 40 <= value <= 400:
                values.append(value)
        return values

    for index in sorted(header_candidates, reverse=True):
        values = candidate_values(index)
        if len(values) >= 45 and 60 <= statistics.median(values) <= 300:
            return index

    # Merged Excel headers vary by edition. Infer the bill column from state
    # rows: customer counts and kWh are much larger, while cents/kWh are much
    # smaller; monthly residential bills cluster roughly in the $60-$300 range.
    ranked: list[tuple[int, float, int]] = []
    for index in range(width):
        if index == state_col:
            continue
        values = candidate_values(index)
        if len(values) < 45:
            continue
        median = statistics.median(values)
        if 60 <= median <= 300:
            ranked.append((len(values), median, index))
    if not ranked:
        raise RuntimeError("Could not locate EIA average-monthly-bill column")
    ranked.sort(reverse=True)
    return ranked[0][2]


def workbook_year(rows: list[list[object]]) -> int | None:
    for row in rows[:20]:
        for value in row:
            text = normalize(value)
            match = re.search(r"\b(20\d{2})\b", text)
            if match:
                year = int(match.group(1))
                if 2000 <= year <= 2100:
                    return year
    return None


def main() -> None:
    raw = download(EIA_BILL_URL)
    workbook = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]

    state_col = find_state_column(rows)
    bill_col = find_bill_column(rows, state_col)
    bills: dict[str, float] = {}

    for row in rows:
        if state_col >= len(row) or bill_col >= len(row):
            continue
        code = state_code(row[state_col])
        if not code:
            continue
        bill = number(row[bill_col])
        if bill is None or bill < 40 or bill > 400:
            continue
        bills[code] = round(bill, 2)

    if len(bills) < 45:
        raise RuntimeError(f"EIA workbook yielded only {len(bills)} state/DC bill values")

    year = workbook_year(rows)
    existing = json.loads(COST_DATA_PATH.read_text())
    electricity = dict(existing.get("electricity", {}))

    previous_period = electricity.get("billPeriod")
    bill_period = str(year) if year else electricity.get("billPeriod", "latest annual")

    # The price escalator carries a specific published bill year forward to the
    # app's price period, so it is only valid for the bill year it was computed
    # against. A newer workbook already reflects the newer prices; reapplying
    # the old escalator on top would inflate them a second time. Reset it here
    # and record the new base period, and the application falls back to using
    # the freshly published bills as-is until a new escalator is derived.
    if bill_period != previous_period:
        electricity["priceInflationMultiplier"] = 1.0
        electricity["priceInflationBasePeriod"] = bill_period
        electricity["pricePeriod"] = bill_period
        print(
            f"[EIA electricity] bill period {previous_period} -> {bill_period}; "
            "price escalator reset to 1.0 pending a new derivation."
        )

    electricity.update(
        {
            "stateAverageMonthlyBill": bills,
            "billPeriod": bill_period,
            "source": "U.S. EIA residential average monthly bill by state",
            "sourceUrl": EIA_BILL_URL,
            # The EIA figure is for an average residential customer, while this
            # app models one renter. The factor is intentionally explicit and
            # editable in generated data rather than hidden in application code.
            "singleRenterFactor": electricity.get("singleRenterFactor", 0.60),
            # Water, trash, gas and home internet stay anchored to the existing
            # metro-specific modeled utility benchmark.
            "modeledNonElectricShare": electricity.get("modeledNonElectricShare", 0.45),
        }
    )
    existing["electricity"] = electricity

    # Keep the eia- segment of the human-readable data version in step with the
    # workbook that was actually parsed. Other segments are owned by their own
    # refresh scripts and are preserved untouched.
    retained = [
        part
        for part in str(existing.get("dataVersion", "")).split("_")
        if part and not part.startswith("eia-")
    ]
    existing["dataVersion"] = "_".join([*retained, f"eia-{bill_period}"])

    existing.setdefault("sources", {})["utilities"] = (
        "Electricity: U.S. EIA residential average monthly bill by state, scaled to a single-renter apartment; "
        "water, gas, trash and home internet remain modeled from the metro benchmark and are inflation-indexed."
    )

    COST_DATA_PATH.write_text(json.dumps(existing, indent=2) + "\n")
    print(f"Updated EIA residential electricity bills for {len(bills)} states/DC ({electricity['billPeriod']}).")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"[EIA electricity] refresh skipped; keeping last-known-good data: {exc}", file=sys.stderr)
        # A source failure must never erase or corrupt the existing generated file.
        # Exit zero so other independent public sources can still refresh.
        sys.exit(0)
